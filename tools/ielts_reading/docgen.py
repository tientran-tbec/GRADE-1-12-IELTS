import re,html,json
from docx import Document
from docx.shared import Pt,Cm,RGBColor,Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH,WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
BLUE=RGBColor(0x1F,0x4E,0x79)
def new_doc():
    d=Document()
    s=d.sections[0]; s.page_width=Emu(7560310); s.page_height=Emu(10692130)
    s.left_margin=s.right_margin=Emu(720090); s.top_margin=s.bottom_margin=Emu(647700)
    st=d.styles['Normal']; st.font.name='Calibri'; st.font.size=Pt(11)
    st.element.rPr.rFonts.set(qn('w:eastAsia'),'Calibri')
    st.paragraph_format.space_after=Pt(4)
    return d
TOK=re.compile(r'(<\/?(?:b|i|u|sup|sub|strong|em)>|<br\s*/?>)',re.I)
def add_html(par,h,size=None,bold=False,italic=False,color=None):
    st={'b':bold,'i':italic,'u':False,'sup':False,'sub':False}
    h=h or ''
    h=re.sub(r'<(?!/?(?:b|i|u|sup|sub|strong|em|br)\b)[^>]*>','',h,flags=re.I)
    for part in TOK.split(h):
        if not part: continue
        m=re.match(r'<(/?)(\w+)',part)
        if m and TOK.fullmatch(part):
            tag=m.group(2).lower(); close=bool(m.group(1))
            if tag=='br': par.add_run().add_break(); continue
            tag={'strong':'b','em':'i'}.get(tag,tag)
            st[tag]=(not close) if tag in st else None
            if tag=='b' and close: st['b']=bold
            if tag=='i' and close: st['i']=italic
            continue
        t=html.unescape(part)
        r=par.add_run(t); r.bold=st['b'] or None; r.italic=st['i'] or None; r.underline=st['u'] or None
        if st['sup']: r.font.superscript=True
        if st['sub']: r.font.subscript=True
        if size: r.font.size=Pt(size)
        if color: r.font.color.rgb=color
def P(doc,h='',size=None,bold=False,italic=False,align=None,indent=None,color=None,after=None,keep=False):
    p=doc.add_paragraph()
    if h: add_html(p,h,size,bold,italic,color)
    if align=='c': p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    if indent: p.paragraph_format.left_indent=Cm(indent)
    if after is not None: p.paragraph_format.space_after=Pt(after)
    if keep: p.paragraph_format.keep_with_next=True
    return p
def pagebreak(doc): doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
def blank(n): return '(%d) ………………………………'%n
def aval(v):
    if isinstance(v,list): 
        out=[]; 
        for x in v:
            x=html.unescape(re.sub('<.*?>','',str(x)))
            if x not in out: out.append(x)
        return ' / '.join(out[:3])
    return str(v)
def table_from_html(doc,th):
    rows=re.findall(r'<tr[^>]*>(.*?)</tr>',th,re.S|re.I)
    cells=[re.findall(r'<t[dh][^>]*>(.*?)</t[dh]>',r,re.S|re.I) for r in rows]
    nc=max((len(c) for c in cells),default=0)
    if not nc: return
    t=doc.add_table(rows=len(cells),cols=nc); t.style='Table Grid'
    for i,row in enumerate(cells):
        for j,c in enumerate(row):
            c=re.sub(r'\{\{(\d+)\}\}',lambda m:blank(int(m.group(1))),c)
            cp=t.cell(i,j).paragraphs[0]; add_html(cp,c,10)
    doc.add_paragraph()
def body_html(doc,body,size=11):
    body=re.sub(r'\{\{(\d+)\}\}',lambda m:blank(int(m.group(1))),body or '')
    pos=0
    for m in re.finditer(r'<table.*?</table>',body,re.S|re.I):
        seg=body[pos:m.start()]; 
        for line in re.split(r'<br\s*/?>',seg):
            if line.strip(): P(doc,line.replace('&nbsp;',' '),size)
        table_from_html(doc,m.group(0)); pos=m.end()
    for line in re.split(r'<br\s*/?>',body[pos:]):
        if line.strip(): P(doc,line.replace('&nbsp;',' '),size)
def passage(doc,no,pi,d):
    p=d['passages'][pi]; ti=re.sub('<.*?>','',p.get('title') or '')
    P(doc,f"TEST {no} — Passage {pi+1}"+(f": {ti}" if ti else ''),13,bold=True,color=BLUE,keep=True)
    if p.get('instr'): P(doc,p['instr'],10,italic=True)
    for x in p['paras']:
        if x.get('sub'): P(doc,x['html'],11,bold=True)
        elif x.get('label'):
            par=doc.add_paragraph(); r=par.add_run(x['label']+'  '); r.bold=True; add_html(par,x['html'],11)
        else: P(doc,x['html'],11)
def group(doc,g,d,answers=False):
    a,b=g['q']
    P(doc,'Questions %d–%d'%(a,b) if a!=b else 'Question %d'%a,12,bold=True,keep=True)
    for l in g.get('lines',[]): P(doc,l,10.5,italic=True,after=1)
    A=d['answers']; t=g['type']
    if t in('tfng','ynng'):
        for it in g['items']:
            P(doc,f"{it['n']}. {it['q']}",bold=True,keep=True)
            P(doc,' / '.join(g['labels'])+(f"      →  {aval(A[str(it['n'])])}" if answers else ''),10)
    elif t=='mcq':
        for it in g['items']:
            P(doc,f"{it['n']}. {it['q']}",bold=True,keep=True,after=1)
            for op in it['opts']: P(doc,f"{op['k']}  {op.get('t','')}",10.5,indent=0.6,after=0)
            if answers: P(doc,f"→ {aval(A[str(it['n'])])}",10,italic=True)
    elif t=='mcq_multi':
        for it in g['items']:
            ns=it['ns']; P(doc,f"Q{ns[0]}-{ns[-1]}. {it['q']}" if len(ns)>1 else f"{ns[0]}. {it['q']}",bold=True,keep=True,after=1)
            for op in it['opts']: P(doc,f"{op['k']}  {op.get('t','')}",10.5,indent=0.6,after=0)
            if answers: P(doc,"→ "+', '.join(aval(A[str(n)]) for n in ns),10,italic=True)
    elif t=='match':
        if g.get('options'):
            P(doc,g.get('options_title') or 'List',10.5,bold=True,after=1,keep=True)
            for op in g['options']: P(doc,f"{op['k']}.  {op['t']}",10.5,indent=0.6,after=0)
        for it in g['items']:
            q=re.sub(r'\{\{%d\}\}'%it['n'],'…………',it.get('q',''))
            P(doc,f"{it['n']}. {q}   →  "+(aval(A[str(it['n'])]) if answers else 'Đáp án: …………………………'),bold=True)
    elif t=='completion':
        if g.get('bank'):
            P(doc,'Word / phrase list',10.5,bold=True,after=1,keep=True)
            for x in g['bank']: P(doc,f"{x['k']}  {x['t']}",10.5,indent=0.6,after=0)
        if g.get('title'): P(doc,g['title'],12,bold=True,keep=True)
        body_html(doc,g.get('body',''))
    else:
        body_html(doc,g.get('body',''))
def key_rows(d,types=None):
    rows=[]
    for g in d['groups']:
        if types and g['type'] not in types: continue
        if g['type']=='completion': ns=[int(x) for x in re.findall(r'\{\{(\d+)\}\}',g['body'])]
        elif g['type']=='mcq_multi': ns=[n for it in g['items'] for n in it['ns']]
        else: ns=[it['n'] for it in g['items']]
        for n in ns: rows.append((n,aval(d['answers'][str(n)])))
    return rows
def shade(cell,color):
    tcPr=cell._tc.get_or_add_tcPr(); sh=OxmlElement('w:shd'); sh.set(qn('w:val'),'clear'); sh.set(qn('w:color'),'auto'); sh.set(qn('w:fill'),color); tcPr.append(sh)
def key_table(doc,rows,with_test=False):
    t=doc.add_table(rows=1,cols=3 if with_test else 2); t.style='Table Grid'; t.alignment=WD_TABLE_ALIGNMENT.CENTER
    hdr=(['Test','Câu hỏi','Đáp án đúng'] if with_test else ['Câu hỏi','Đáp án đúng'])
    for i,h in enumerate(hdr):
        c=t.rows[0].cells[i]; c.text=''; r=c.paragraphs[0].add_run(h); r.bold=True; r.font.size=Pt(10); shade(c,'D9E2F3')
    for row in rows:
        cs=t.add_row().cells
        for i,v in enumerate(row):
            cs[i].text=''; r=cs[i].paragraphs[0].add_run(str(v)); r.font.size=Pt(10)
def add_test(doc,d,answers_key=True):
    no=d['no']
    for pi in range(3):
        passage(doc,no,pi,d)
        for g in [g for g in d['groups'] if g['p']==pi+1]: group(doc,g,d)
        P(doc,'')
    if answers_key:
        P(doc,f'TEST {no} — ĐÁP ÁN / ANSWER KEY',13,bold=True,color=BLUE,keep=True)
        key_table(doc,key_rows(d)); P(doc,'')
